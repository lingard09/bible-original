"""기본 뜻 검토 목록(review/기본뜻-검토목록.xlsx)을 사이트 데이터에 반영합니다.

사용법: python tools/apply_gloss_review.py review/기본뜻-검토목록.xlsx [--data data] [--dry-run]

- 검토 결과가 '수정'이고 '수정한 뜻'이 있으면 그 뜻으로 바꿉니다.
  (data/glosses-he.json, data/glosses-gr.json, data/vocab-*.json, data/quiz-*.json, data/alphabet.json)
- '확인' 또는 '수정'인 단어는 data/gloss-review.json에 기록해 사이트에서 '검토한 뜻'으로 표시합니다.
- '보류'와 빈칸은 건드리지 않습니다.
"""
import argparse, json, os, re, sys
from openpyxl import load_workbook

SHEETS = {'히브리어': 'he', '헬라어': 'gr'}


def load(path):
    with open(path, encoding='utf8') as f:
        return json.load(f)


def save(path, data):
    with open(path, 'w', encoding='utf8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))


def read_review(path):
    wb = load_workbook(path, data_only=True)
    out = {'he': [], 'gr': []}
    problems = []
    for name, lang in SHEETS.items():
        ws = wb[name]
        header = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(header)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw_id, result = row[col['ID']], (row[col['검토 결과']] or '').strip()
            if not raw_id or result not in ('확인', '수정'):
                continue
            # 히브리어 ID는 'H559', 헬라어 ID는 '0 (G3588)'처럼 앞의 숫자가 사이트 표제어 번호
            wid = str(raw_id).strip() if lang == 'he' else re.match(r'\d+', str(raw_id)).group()
            new = (row[col['수정한 뜻']] or '').strip()
            if result == '수정' and not new:
                problems.append(f'{name} {row[col["표제어"]]}: 수정인데 수정한 뜻이 비어 있어 건너뜁니다.')
                continue
            out[lang].append((wid, result, new, row[col['표제어']], row[col['현재 기본 뜻']]))
    return out, problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--data', default=os.path.join(os.path.dirname(__file__), '..', 'data'))
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()
    d = a.data
    review, problems = read_review(a.xlsx)
    glosses = {'he': load(f'{d}/glosses-he.json'), 'gr': load(f'{d}/glosses-gr.json')}
    vocab = {l: load(f'{d}/vocab-{l}.json') for l in ('he', 'gr')}
    quiz = {l: load(f'{d}/quiz-{l}.json') for l in ('he', 'gr')}
    alphabet = load(f'{d}/alphabet.json')
    state_path = f'{d}/gloss-review.json'
    state = load(state_path) if os.path.exists(state_path) else {'he': {}, 'gr': {}}
    changed = {'he': 0, 'gr': 0}
    for lang, rows in review.items():
        vmap = {w['id']: w for w in vocab[lang]['words']}
        for wid, result, new, lemma, old in rows:
            if wid not in vmap:
                problems.append(f'{lemma} ({wid}): 사이트 단어 목록에 없는 ID라 건너뜁니다.')
                continue
            if result == '수정' and new != vmap[wid]['g']:
                if lang == 'he':
                    glosses['he'][wid[1:]] = new
                else:
                    glosses['gr']['glosses'][int(wid)] = new
                vmap[wid]['g'] = new
                if wid in quiz[lang]['lemmas']:
                    quiz[lang]['lemmas'][wid][1] = new
                # 원어 입문 예시: 히브리어는 Strong 번호, 헬라어는 표제어로 같은 단어를 찾음
                for group in (alphabet['he']['letters'], alphabet['he']['finals'], alphabet['he']['vowels']) if lang == 'he' else (alphabet['gr']['letters'],):
                    for item in group.values():
                        if (lang == 'he' and item.get('strong') == wid) or (lang == 'gr' and item.get('lemma') == lemma):
                            item['gloss'] = new
                for item in alphabet['practice'][lang]:
                    if (lang == 'he' and item.get('strong') == wid) or (lang == 'gr' and item.get('lemma') == lemma):
                        item['gloss'] = new
                changed[lang] += 1
                print(f'수정 {lemma}: {old} → {new}')
            state[lang][wid] = 'edited' if result == '수정' else 'ok'
    for p in problems:
        print('주의:', p, file=sys.stderr)
    print(f"히브리어 검토 {len(state['he'])}개(이번에 뜻 수정 {changed['he']}개), 헬라어 검토 {len(state['gr'])}개(이번에 뜻 수정 {changed['gr']}개)")
    if a.dry_run:
        print('--dry-run: 파일을 저장하지 않았습니다.')
        return
    save(f'{d}/glosses-he.json', glosses['he'])
    save(f'{d}/glosses-gr.json', glosses['gr'])
    for l in ('he', 'gr'):
        save(f'{d}/vocab-{l}.json', vocab[l])
        save(f'{d}/quiz-{l}.json', quiz[l])
    save(f'{d}/alphabet.json', alphabet)
    save(state_path, state)


if __name__ == '__main__':
    main()
